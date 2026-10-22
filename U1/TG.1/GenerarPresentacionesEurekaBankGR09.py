from __future__ import annotations

import argparse
import hashlib
import os
import posixpath
import re
import shutil
import sys
import tempfile
import textwrap
import zipfile
from pathlib import Path, PurePosixPath
from xml.etree import ElementTree as ET

LIBS = Path(r"D:\tools\eureka-pptx-libs")
if LIBS.is_dir():
    sys.path.insert(0, str(LIBS))

from PIL import Image, ImageDraw, ImageFont, ImageOps
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_AUTO_SIZE


TEMPLATE = Path("EUREKABANK_SOAP_JAVA_ANDRANGO_CALAPAQUI_QUILUMBAQUIN") / "04. DOCUMENTACIÓN" / "SOAP_EUREKABANK_BDD_LOCAL_JAVA.pptx"
SLIDE_W, SLIDE_H = 12192000, 6858000
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
OFFICE_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PPT_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
EMU = 914400

PROFILES = [
    {"folder": "1.1. PAO202651_SOAP_JAVA_CONBDD_NRC36884_GR09", "key": "SOAP_JAVA", "stack": "Java", "protocol": "SOAP", "db": "MySQL", "server": "Jakarta EE · JAX-WS · Payara Server Community 7", "endpoint": "http://localhost:8080/WS_EurekaBank_SOAP_JAVA_GR09/WSEureka", "endpoint_env": "EUREKABANK_SOAP_URL", "dbvars": "EUREKABANK_DB_URL · EUREKABANK_DB_USER · EUREKABANK_DB_PASSWORD", "clients": "consola Java · escritorio Swing · web JSP/Servlet · Android"},
    {"folder": "1.2. PAO202651_SOAP_DOTNET_CONBDD_NRC36884_GR09", "key": "SOAP_DOTNET", "stack": ".NET", "protocol": "SOAP", "db": "SQL Server", "server": ".NET 10 · CoreWCF", "endpoint": "http://localhost:5080/WSEureka", "endpoint_env": "EUREKABANK_SOAP_DOTNET_URL", "dbvars": "EUREKABANK_DB_CONNECTION_STRING", "clients": "consola .NET · WPF · ASP.NET Core MVC · .NET MAUI"},
    {"folder": "1.3. PAO202651_RESTFUL_JAVA_CONBDD_NRC36884_GR09", "key": "RESTFUL_JAVA", "stack": "Java", "protocol": "REST", "db": "MySQL", "server": "Jakarta EE · Jakarta REST · Payara Server Community 7", "endpoint": "http://127.0.0.1:8080/eurekabank-rest-java-gr09/api/eureka", "endpoint_env": "EUREKABANK_REST_URL", "dbvars": "EUREKABANK_DB_URL · EUREKABANK_DB_USER · EUREKABANK_DB_PASSWORD", "clients": "consola Java · escritorio Swing · web JSP/Servlet · Android"},
    {"folder": "1.4. PAO202651_RESTFUL_DOTNET_CONBDD_NRC36884_GR09", "key": "RESTFUL_DOTNET", "stack": ".NET", "protocol": "REST", "db": "SQL Server", "server": ".NET 10 · ASP.NET Core", "endpoint": "http://localhost:5081/api/eureka", "endpoint_env": "EUREKABANK_REST_DOTNET_URL", "dbvars": "EUREKABANK_DB_CONNECTION_STRING", "clients": "consola .NET · WPF · ASP.NET Core MVC · .NET MAUI"},
]

def set_paragraph_text(paragraph, text: str) -> None:
    runs = paragraph.runs
    if not runs:
        paragraph.add_run().text = text
        return
    runs[0].text = text
    for run in runs[1:]:
        run.text = ""


def set_shape_lines(shape, lines: list[str]) -> None:
    paragraphs = shape.text_frame.paragraphs
    if len(lines) > len(paragraphs):
        raise ValueError(f"Texto no cabe en la plantilla: {shape.name} tiene {len(paragraphs)} párrafos y recibió {len(lines)}")
    shape.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    for index, paragraph in enumerate(paragraphs):
        set_paragraph_text(paragraph, lines[index] if index < len(lines) else "")


def replace_shape_paragraphs(shape, transform) -> None:
    for paragraph in shape.text_frame.paragraphs:
        text = paragraph.text
        updated = transform(text)
        if updated != text:
            set_paragraph_text(paragraph, updated)


def get_shape(slide, shape_id: int):
    return next((shape for shape in slide.shapes if shape.shape_id == shape_id), None)


def diagram_assets(project: Path) -> dict[str, Path]:
    uml = project / "01. UML" / "03. DIAGRAMAS UML"
    mer = project / "02. MER"

    def one(folder: Path, pattern: str, exclude_detail: bool = False) -> Path:
        matches = list(folder.glob(pattern))
        if exclude_detail:
            matches = [path for path in matches if not re.search(r"_(?:02|03|04|05|06|07)_", path.stem)]
        if len(matches) != 1:
            raise FileNotFoundError(f"Se esperaba una imagen para {pattern}; encontradas: {len(matches)}")
        return matches[0]

    cases = uml / "01_CASOS_DE_USO"
    activities = uml / "02_ACTIVIDAD"
    sequences = uml / "03_SECUENCIA"
    classes = uml / "04_CLASES"
    architecture = uml / "05_ARQUITECTURA"
    conceptual = mer / "01. CONCEPTUAL"
    logical = mer / "02. LOGICO"
    physical = mer / "03. FISICO"
    return {
        "architecture": one(architecture, "GR09_Arquitectura_*.png"),
        "class": one(classes, "GR09_Clases_*.png"),
        "case_summary": one(cases, "GR09_CasosUso_CU_Resumen.png"),
        "case_login": one(cases, "GR09_CasosUso_CU_Login.png"),
        "case_movements": one(cases, "GR09_CasosUso_CU_Movimientos.png"),
        "case_deposit": one(cases, "GR09_CasosUso_CU_Deposito.png"),
        "case_transfer": one(cases, "GR09_CasosUso_CU_Transferencia.png"),
        "activity_login": one(activities, "GR09_Actividad_DA_Login.png"),
        "activity_movements": one(activities, "GR09_Actividad_DA_Movimientos.png"),
        "activity_deposit": one(activities, "GR09_Actividad_DA_Deposito.png"),
        "activity_transfer": one(activities, "GR09_Actividad_DA_Transferencia.png"),
        "sequence_login": one(sequences, "GR09_Secuencia_DS_Login.png"),
        "sequence_movements": one(sequences, "GR09_Secuencia_DS_Movimientos.png"),
        "sequence_transfer": one(sequences, "GR09_Secuencia_DS_Transferencia.png"),
        "conceptual": one(conceptual, "GR09_EurekaBank_Conceptual_*.png", True),
        "conceptual_transfer": one(conceptual, "GR09_EurekaBank_Conceptual_*_02_Transferencia_Origen.png"),
        "conceptual_destination": one(conceptual, "GR09_EurekaBank_Conceptual_*_03_Transferencia_Destino.png"),
        "logical": one(logical, "GR09_EurekaBank_Logical_*.png", True),
        "logical_transfer": one(logical, "GR09_EurekaBank_Logical_*_02_Transferencia_Origen.png"),
        "logical_destination": one(logical, "GR09_EurekaBank_Logical_*_03_Transferencia_Destino.png"),
        "logical_movements": one(logical, "GR09_EurekaBank_Logical_*_04_Movimientos.png"),
        "physical": one(physical, "GR09_EurekaBank_Physical_*.png", True),
        "physical_transfer": one(physical, "GR09_EurekaBank_Physical_*_02_Transferencia_Origen.png"),
        "physical_destination": one(physical, "GR09_EurekaBank_Physical_*_03_Transferencia_Destino.png"),
        "physical_movements": one(physical, "GR09_EurekaBank_Physical_*_04_Movimientos.png"),
        "interface": one(project / "05. DOCUMENTACION" / "01. MANUAL TECNICO" / "Capturas", "InterfazInicioSesion.png"),
    }


def slide_text(profile: dict, number: int) -> list[str] | None:
    if number == 3:
        return ["2.1 OBJETIVO", "", f"Documentar EurekaBank con {profile['protocol']} y {profile['stack']}: acceso, movimientos, depósitos y transferencias. El servicio valida las reglas bancarias; cuatro clientes comparten el mismo contrato."]
    if number == 4:
        return ["2.2 MODELO CONCEPTUAL DE LA BASE DE DATOS", "El modelo vincula titular, cuentas y movimientos. Transferencia conserva origen y destino y genera débito y crédito."]
    if number == 5:
        return ["2.3 PREPARAR LA BASE DE DATOS", f"Cree el esquema de EurekaBank en {profile['db']}. Mantenga usuario y contraseña en variables de entorno."]
    if number == 6:
        return ["Defina tablas, claves y relaciones. Las 19 entidades incluyen cuentas, movimientos y Transferencia."]
    if number == 7:
        return ["Revise claves, estados, saldos y unicidad de la clave idempotente antes de habilitar operaciones."]
    if number == 8:
        return ["Use datos autorizados. No incluya contraseñas, tokens ni cadenas de conexión en capturas o material compartido."]
    if number == 9:
        return [f"2.4 SERVICIO {profile['protocol']}", f"{profile['server']}. El servidor concentra autenticación, reglas y persistencia en {profile['db']}.", f"Endpoint: {profile['endpoint']}", ""]
    if number == 10:
        return [f"Servidor {profile['server']} · endpoint {profile['endpoint']} · base {profile['db']}."]
    if number == 11:
        return [f"Variables solo en servidor: {profile['dbvars']}. Nunca las incluya en clientes ni capturas."]
    if number == 12:
        return [f"Confirme contrato o rutas, puerto, conexión a {profile['db']} y esquema actualizado antes de probar clientes."]
    if number == 13:
        architecture = "Paquetes db, modelos, vistas, controladores, servicios y ws siguen ec.edu.monster." if profile["stack"] == "Java" else "Espacios de nombres .NET separan controladores, servicios, modelos y repositorios."
        return ["5. Separar responsabilidades", "MVC mantiene vista, controlador y modelo; persistencia queda en servidor.", architecture]
    if number == 14:
        return ["Movimiento registra cuenta, fecha, importe y tipo. Transferencia enlaza los dos apuntes contables de una operación."]
    if number == 15:
        return ["Repositorio usa consultas parametrizadas y libera conexión, comandos y lectores al terminar."]
    if number == 16:
        return ["Valide titularidad, cuentas activas, moneda, saldo e importe. Confirme juntos débito 009 y crédito 008."]
    if number == 17:
        build = "mvn test · mvn package · WAR en Payara 7" if profile["stack"] == "Java" else "dotnet build · dotnet test · servicio .NET 10"
        return [f"10. Prepare {build}; inicie servicio y configure {profile['endpoint']}. Pruebas requieren servidor y base activos."]
    if number == 18:
        return ["2.5 CLIENTE DE CONSOLA", f"La consola {profile['stack']} usa MVC, consulta el servicio {profile['protocol']} y presenta sesión, movimientos, depósitos y transferencias."]
    if number == 19:
        transport = "WSDL y mensajes XML" if profile["protocol"] == "SOAP" else "HTTP, JSON y token Bearer"
        return [f"Use {transport}; endpoint configurable: {profile['endpoint']}. El fallo de red muestra un mensaje útil y permite reintentar."]
    if number == 21:
        return ["3. MVC: vista presenta datos; controlador valida interacción; modelo conserva estado mínimo. El servidor aplica reglas bancarias."]
    if number == 22:
        return [f"El adaptador encapsula {profile['protocol']}; la sesión conserva token temporal y borra contraseña tras autenticar."]
    if number == 23:
        return ["Valide cuenta de ocho dígitos e importe positivo con máximo dos decimales. Rechace origen igual a destino."]
    if number == 24:
        return ["El menú ofrece movimientos, depósito y transferencias a cuentas propias o de terceros."]
    if number == 25:
        return ["Depósito acredita una cuenta. Transferencia debita origen y acredita destino en una transacción."]
    if number == 26:
        desktop = "Swing" if profile["stack"] == "Java" else "WPF"
        return [f"2.6 CLIENTE DE ESCRITORIO · {desktop}: consume el servicio; no guarda credenciales ni accede a {profile['db']}."]
    if number == 27:
        return ["Use endpoint configurable y token temporal; borre credenciales al cerrar sesión."]
    if number == 28:
        transport = "proxy JAX-WS" if profile["protocol"] == "SOAP" and profile["stack"] == "Java" else "cliente CoreWCF" if profile["protocol"] == "SOAP" else "HttpClient y JSON"
        return [f"El cliente usa {transport}, valida respuesta y libera recursos al finalizar la operación."]
    if number == 29:
        return ["Rechace campos vacíos, importe inválido, titular distinto, moneda diferente y saldo insuficiente."]
    if number == 30:
        return ["Confirme éxito tras respuesta del servicio. Reintente respuesta incierta con la misma clave idempotente."]
    if number == 31:
        return ["6. Inicio de sesión: captura del formulario vacío; no acredita conexión ni autenticación exitosa."]
    if number == 32:
        return ["7. Movimientos: muestre fecha, tipo e importe; 009 SALIDA es débito y 008 INGRESO es crédito."]
    if number == 33:
        return ["8. Transferir: seleccione origen y destino, valide titularidad y fondos, y conserve la clave idempotente."]
    if number == 34:
        web = "JSP/Servlet" if profile["stack"] == "Java" else "ASP.NET Core MVC"
        return [f"2.7 CLIENTE WEB · {web}", f"Protege sesión y consume el contrato {profile['protocol']} compartido por las otras interfaces."]
    if number == 35:
        web = "Servlets y JSP" if profile["stack"] == "Java" else "controladores, vistas y servicios ASP.NET Core MVC"
        return [f"1. Capas web: {web}; el cliente no conecta directo a {profile['db']}."]
    if number == 36:
        return ["2. Acceso web: proteja sesión, no exponga tokens y vuelva a login al revocar sesión."]
    if number == 37:
        return ["3. Consultas: muestre movimientos de cuenta autorizada; no acepte titular manipulado."]
    if number == 38:
        return ["4. Operaciones: valide el formulario y muestre el resultado que confirma el servicio."]
    if number == 39:
        mobile = "Android" if profile["stack"] == "Java" else ".NET MAUI"
        return [f"1. CLIENTE MÓVIL · {mobile}: configure endpoint accesible y conserve solo token temporal."]
    if number == 40:
        return ["2. MVC: vista captura, controlador coordina cliente y modelo presenta respuesta; no persista secretos."]
    if number == 41:
        return ["3. Red móvil: localhost refiere al dispositivo; use una dirección alcanzable desde emulador o teléfono."]
    if number == 42:
        if profile["protocol"] == "SOAP":
            return [f"EurekaBank expone un contrato SOAP/WSDL con {profile['server']} y persistencia en {profile['db']}. Consola, escritorio, web y móvil consumen las mismas operaciones.", "La transferencia valida titularidad, estado, moneda, saldo e importe. Registra 009 SALIDA y 008 INGRESO en una transacción; un error revierte ambos apuntes."]
        return [f"EurekaBank expone recursos REST/JSON con {profile['server']} y persistencia en {profile['db']}. Consola, escritorio, web y móvil comparten rutas y sesión temporal.", "La transferencia valida titularidad, estado, moneda, saldo e importe. Registra 009 SALIDA y 008 INGRESO en una transacción; un error revierte ambos apuntes."]
    if number == 43:
        build = "Payara 7 y Maven" if profile["stack"] == "Java" else ".NET 10 y dotnet"
        return [f"Configure {build} y {profile['db']} antes de abrir clientes. Use variables de entorno para endpoint y conexión; no copie secretos en capturas.", "Pruebe titularidad, moneda, saldo insuficiente, doble solicitud y rollback. Mantenga la misma clave idempotente al resolver una respuesta incierta.", "En producción use HTTPS, permisos mínimos y registros sin contraseñas ni tokens completos."]
    return None


def image_roles(number: int) -> list[str]:
    return {
        4: ["conceptual"],
        5: ["db_setup", "transfer_schema"], 6: ["conceptual_destination"],
        7: ["physical", "activity_login"], 8: ["physical_transfer"],
        9: ["conceptual", "sequence_login"], 10: ["case_summary"], 11: ["sequence_login"],
        12: ["logical"], 13: ["architecture", "conceptual"],
        14: ["logical_movements"], 15: ["sequence_movements"], 16: ["logical_transfer"],
        17: ["architecture"], 18: ["interface", "case_summary"], 19: ["case_login", "sequence_login"],
        20: ["conceptual_transfer", "interface"], 21: ["logical"], 22: ["sequence_login", "interface"],
        23: ["logical_transfer", "logical_movements"], 24: ["case_summary"], 25: ["sequence_movements"],
        26: ["interface"], 27: ["sequence_login", "sequence_movements"],
        28: ["sequence_transfer", "interface"], 29: ["logical_transfer", "conceptual"],
        30: ["sequence_transfer", "architecture", "logical_transfer"], 31: ["interface"],
        32: ["sequence_movements"], 33: ["sequence_transfer"],
        34: ["sequence_login", "interface"], 35: ["interface"], 36: ["sequence_login", "interface"],
        37: ["sequence_movements", "case_movements"], 38: ["sequence_login", "logical_movements"],
        39: ["mobile_endpoint", "case_movements", "conceptual"],
        40: ["logical_transfer", "architecture", "case_transfer", "case_login", "conceptual"],
        41: ["activity_login", "sequence_transfer", "activity_transfer"],
    }.get(number, ["architecture"])


def make_code_capture(title: str, lines: list[str], output: Path, ratio: float) -> Path:
    width = 1800
    height = max(230, round(width / ratio))
    image = Image.new("RGB", (width, height), (25, 28, 34))
    draw = ImageDraw.Draw(image)
    code_path = Path(r"C:\Windows\Fonts\consola.ttf")
    title_path = Path(r"C:\Windows\Fonts\arialbd.ttf")
    code_size, line_height = (54, 62) if ratio > 4 else (42, 50)
    code_font = ImageFont.truetype(str(code_path), code_size) if code_path.exists() else ImageFont.load_default()
    title_font = ImageFont.truetype(str(title_path), 33) if title_path.exists() else ImageFont.load_default()
    draw.rectangle((14, 14, width - 14, height - 14), outline=(94, 110, 132), width=2)
    draw.text((34, 28), title, fill=(146, 207, 255), font=title_font)
    code_lines = [part for line in lines for part in textwrap.wrap(line, 72, subsequent_indent="    ") or [line]]
    if len(code_lines) > max(1, (height - 78) // line_height):
        raise ValueError(f"La captura de código no cabe: {title}")
    for index, line in enumerate(code_lines):
        draw.text((34, 70 + index * line_height), line, fill=(232, 235, 240), font=code_font)
    image.save(output, format="PNG", optimize=True)
    return output


def sql_captures(project: Path, profile: dict, output_dir: Path) -> dict[str, Path]:
    migration_name = "Migrar_Transferencias_MySQL.sql" if profile["db"] == "MySQL" else "Migrar_Transferencias_SQLServer.sql"
    migration = (project / "03. BDD" / migration_name).read_text(encoding="utf-8-sig")
    if profile["db"] == "MySQL":
        anchor = next(line.strip() for line in migration.splitlines() if line.strip().upper().startswith("USE "))
    else:
        anchor = next(line.strip() for line in migration.splitlines() if line.strip().upper().startswith("SET XACT_ABORT"))
    table = re.search(r"(?is)CREATE\s+TABLE\s+(?:dbo\.)?Transferencias\s*\((.*?)\n\s*\)", migration)
    if table is None or "clave_operacion" not in table.group(1).casefold() or "unique" not in table.group(1).casefold():
        raise ValueError(f"No se encontró DDL de Transferencias para {profile['key']}")
    fields = [line.strip() for line in table.group(1).splitlines() if line.strip()]
    key_fields = [
        line
        for line in fields
        if re.match(r"(?i)(id_transferencia|clave_operacion|chr_cuencodigo_origen|chr_cuencodigo_destino|dec_importe)\s", line)
        or "uq_transferencias_usuario_clave" in line.casefold()
        or "ck_transferencias_cuentas" in line.casefold()
    ]
    if len(key_fields) < 7:
        raise ValueError(f"El esquema de transferencias no contiene todos los campos clave: {profile['key']}")
    return {
        "db_setup": make_code_capture(
            f"EUREKABANK · {profile['db']} · GRUPO 09", [anchor], output_dir / "database-setup.png", 6.62
        ),
        "transfer_schema": make_code_capture(
            f"TRANSFERENCIAS · {profile['db']} · EXTRACTO DE CAMPOS CLAVE",
            key_fields,
            output_dir / "transfer-schema.png",
            2.7,
        ),
        "mobile_endpoint": make_code_capture(
            "CLIENTE MÓVIL · CONFIGURE UNA DIRECCIÓN ALCANZABLE",
            [f"{profile['endpoint_env']}=<URL accesible desde el dispositivo>"],
            output_dir / "mobile-endpoint.png",
            10.18,
        ),
    }


def frame_image(source: Path, ratio: float, output: Path) -> Path:
    width = 1800
    height = max(220, round(width / ratio))
    with Image.open(source) as opened:
        original = opened.convert("RGB")
        contained = ImageOps.contain(original, (width - 30, height - 30), method=Image.Resampling.LANCZOS)
        canvas = Image.new("RGB", (width, height), "white")
        canvas.paste(contained, ((width - contained.width) // 2, (height - contained.height) // 2))
        canvas.save(output, format="PNG", optimize=True)
    return output


def set_picture(picture, image_path: Path, description: str) -> None:
    _, rid = picture.part.get_or_add_image_part(str(image_path))
    picture._element.blipFill.blip.rEmbed = rid
    picture.crop_left = picture.crop_right = picture.crop_top = picture.crop_bottom = 0
    props = picture._element.nvPicPr.cNvPr
    props.set("descr", description)
    props.set("title", description)


def prune_unused_media(path: Path) -> None:
    ET.register_namespace("", REL_NS)
    ET.register_namespace("r", OFFICE_REL_NS)
    with zipfile.ZipFile(path, "r") as source:
        entries = {item.filename: (item, source.read(item.filename)) for item in source.infolist()}

    for rel_name, (info, payload) in list(entries.items()):
        if not rel_name.endswith(".rels") or "/_rels/" not in rel_name:
            continue
        rel_path = PurePosixPath(rel_name)
        source_part = str(rel_path.parent.parent / rel_path.name[:-5])
        source_xml = entries.get(source_part)
        if source_xml is None:
            continue
        xml_root = ET.fromstring(source_xml[1])
        used_ids = set()
        for element in xml_root.iter():
            for attr_name, value in element.attrib.items():
                if attr_name.startswith("{" + OFFICE_REL_NS + "}"):
                    used_ids.add(value)
        rel_root = ET.fromstring(payload)
        changed = False
        for relationship in list(rel_root):
            if relationship.attrib.get("Type", "").endswith("/image") and relationship.attrib.get("Id") not in used_ids:
                rel_root.remove(relationship)
                changed = True
        if changed:
            entries[rel_name] = (info, ET.tostring(rel_root, encoding="utf-8", xml_declaration=True))

    targets = set()
    for rel_name, (_, payload) in entries.items():
        if not rel_name.endswith(".rels"):
            continue
        rel_root = ET.fromstring(payload)
        rel_path = PurePosixPath(rel_name)
        source_part = str(rel_path.parent.parent / rel_path.name[:-5]) if "/_rels/" in rel_name else ""
        source_dir = posixpath.dirname(source_part)
        for relationship in rel_root:
            target = relationship.attrib.get("Target", "")
            if not target or relationship.attrib.get("TargetMode") == "External":
                continue
            resolved = posixpath.normpath(target.lstrip("/") if target.startswith("/") else posixpath.join(source_dir, target))
            targets.add(resolved)
    for name in list(entries):
        if name.startswith("ppt/media/") and name not in targets:
            del entries[name]

    temporary = path.with_suffix(path.suffix + ".cleaning")
    try:
        with zipfile.ZipFile(temporary, "w", compression=zipfile.ZIP_DEFLATED) as output:
            for name, (info, payload) in entries.items():
                output.writestr(info, payload)
        temporary.replace(path)
    finally:
        if temporary.exists():
            temporary.unlink()


def update_deck(root: Path, profile: dict, template: Path) -> Path:
    project = root / profile["folder"]
    output = project / "05. DOCUMENTACION" / "02. PRESENTACION" / f"PRESENTACION_EUREKABANK_{profile['key']}_GR09.pptx"
    output.parent.mkdir(parents=True, exist_ok=True)
    deck = Presentation(str(template))
    if len(deck.slides) != 44 or deck.slide_width != SLIDE_W or deck.slide_height != SLIDE_H:
        raise ValueError("La plantilla original cambió de estructura o formato 16:9")

    assets = diagram_assets(project)
    nav_protocol = "SOAP" if profile["protocol"] == "SOAP" else "REST"
    nav = ["MARCO TEÓRICO", "DESARROLLO", "OBJETIVO", "MODELO CONCEPTUAL DE LA BASE DE DATOS", "CREAR BASE DE DATOS", f"CREACIÓN DEL SERVIDOR {nav_protocol}", "CREACIÓN DEL PROYECTO – CLIENTE CONSOLA", "CREACIÓN DEL PROYECTO – CLIENTE ESCRITORIO", "CREACIÓN DEL PROYECTO – CLIENTE WEB", "CREACIÓN DEL PROYECTO – CLIENTE MÓVIL", "CONCLUSIONES", "RECOMENDACIONES", "BIBLIOGRAFÍA"]
    toc = ["DESARROLLO", "OBJETIVO", "MODELO CONCEPTUAL Y LÓGICO", "MODELO FÍSICO Y TRANSFERENCIAS", "CONFIGURAR BASE DE DATOS", "APLICAR ESQUEMA Y RESTRICCIONES", f"SERVICIO {profile['protocol']}", "PERSISTENCIA Y SEGURIDAD", "CLIENTE DE CONSOLA", "CLIENTE DE ESCRITORIO", "CLIENTE WEB", "CLIENTE MÓVIL", "CONCLUSIONES Y RECOMENDACIONES", "BIBLIOGRAFÍA"]
    with tempfile.TemporaryDirectory(prefix="eureka-pptx-", dir=r"D:\tools") as temp_dir:
        temp = Path(temp_dir)
        code_captures = sql_captures(project, profile, temp)
        for number, slide in enumerate(deck.slides, 1):
            if number == 1:
                title = get_shape(slide, 2)
                if title and title.has_text_frame:
                    paragraph = title.text_frame.paragraphs[0]
                    if len(paragraph.runs) > 1:
                        paragraph.runs[0].text = "MANUAL"
                        paragraph.runs[1].text = f"EUREKABANK CON {profile['protocol']} · {profile['db'].upper()} · {profile['stack'].upper()}"
                    else:
                        set_shape_lines(title, [f"MANUAL EUREKABANK · {profile['protocol']} · {profile['stack']}"])
                group = get_shape(slide, 3)
                if group:
                    set_shape_lines(group, ["Grupo 09 · NRC 36884", "Arquitectura de software", "EurekaBank · 2026", ""])
                    group.top = int(5.26 * EMU)
                date = get_shape(slide, 5)
                if date:
                    set_shape_lines(date, ["FECHA: 22/10/2026", f"VARIANTE: {profile['stack']} · {profile['protocol']}"])
                    date.left = int(8.6 * EMU)
                    date.width = int(4.3 * EMU)
            elif number == 2:
                toc_shape = get_shape(slide, 3)
                if toc_shape:
                    set_shape_lines(toc_shape, toc)
            elif number == 44:
                title = get_shape(slide, 2)
                if title:
                    set_shape_lines(title, ["Bibliografía"])
                refs = get_shape(slide, 3)
                if refs:
                    set_shape_lines(refs, ["[1] EurekaBank. Manual de implementación del servicio SOAP Java y base de datos local. Documento de referencia.", "[2] EurekaBank. Presentación de diseño conceptual y desarrollo del servicio SOAP Java. Documento de referencia."])
            elif number in (42, 43):
                title = get_shape(slide, 2)
                if title:
                    set_shape_lines(title, ["Conclusiones" if number == 42 else "Recomendaciones"])
                body = get_shape(slide, 3)
                if body:
                    lines = slide_text(profile, number)
                    set_shape_lines(body, lines or [])
            else:
                lines = slide_text(profile, number)
                if lines:
                    body = get_shape(slide, 6) if number in (37, 38, 39, 40, 41) else get_shape(slide, 3)
                    if body:
                        set_shape_lines(body, lines)

            for shape in slide.shapes:
                if shape.has_text_frame and number != 1:
                    if shape.shape_id == 5 and number >= 3:
                        set_shape_lines(shape, nav)
                    replace_shape_paragraphs(shape, lambda text: text.replace("CREACIÓN DEL SERVIDOR SOAP", f"CREACIÓN DEL SERVIDOR {nav_protocol}"))
            page_number = get_shape(slide, 4 if number == 1 else 9 if number == 2 else -1)
            if page_number:
                page_number.left -= int(0.14 * EMU)
                page_number.top -= int(0.08 * EMU)

            if number in (1, 2):
                continue
            roles = image_roles(number)
            picture_index = 0
            for shape in list(slide.shapes):
                if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                    continue
                role = roles[min(picture_index, len(roles) - 1)]
                source = code_captures[role] if role in code_captures else assets[role]
                ratio = shape.width / shape.height
                framed = frame_image(source, ratio, temp / f"slide-{number}-picture-{picture_index}.png")
                set_picture(shape, framed, f"EurekaBank Grupo 09 · {profile['stack']} · {profile['protocol']} · {role}")
                picture_index += 1

    properties = deck.core_properties
    properties.title = f"EurekaBank · {profile['protocol']} · {profile['stack']} · Grupo 09"
    properties.subject = f"Arquitectura de software · NRC 36884 · {profile['db']}"
    properties.author = "Grupo 09 · NRC 36884"
    properties.last_modified_by = "Grupo 09"
    properties.keywords = f"EurekaBank, Grupo 09, NRC 36884, {profile['protocol']}, {profile['stack']}, transferencias"
    properties.comments = "Presentación EurekaBank Grupo 09 · NRC 36884."
    deck.save(str(output))
    prune_unused_media(output)
    validate_deck(output, profile, template)
    return output


def validate_deck(path: Path, profile: dict, template: Path) -> None:
    deck = Presentation(str(path))
    if len(deck.slides) != 44 or deck.slide_width != SLIDE_W or deck.slide_height != SLIDE_H:
        raise AssertionError(f"Presentación inválida: {path.name}")
    text = "\n".join(shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)
    required = ["GRUPO 09", "NRC 36884", profile["stack"], profile["protocol"], "Transferencia", "009", "008"]
    missing = [term for term in required if term.casefold() not in text.casefold()]
    if missing:
        raise AssertionError(f"Faltan datos en {path.name}: {missing}")
    if profile["protocol"] == "REST" and "SERVIDOR SOAP" in text:
        raise AssertionError(f"Contenido SOAP en presentación REST: {path.name}")
    if profile["stack"] == "Java" and "WPF" in text:
        raise AssertionError(f"Contenido .NET en presentación Java: {path.name}")
    if profile["stack"] == ".NET" and "ec.edu.monster" in text:
        raise AssertionError(f"Contenido Java en presentación .NET: {path.name}")
    if re.search(r"(?i)\b[A-Z]:\\|/Users/|/home/|U1[/\\]TG\.1", text):
        raise AssertionError(f"La presentación contiene una ruta local: {path.name}")
    active_pic_count = 0
    for slide in deck.slides:
        for shape in slide.shapes:
            if shape.left < 0 or shape.top < 0 or shape.left + shape.width > deck.slide_width + 2 or shape.top + shape.height > deck.slide_height + 2:
                raise AssertionError(f"Elemento fuera de diapositiva en {path.name}")
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                active_pic_count += 1
    if not active_pic_count:
        raise AssertionError(f"La presentación perdió sus imágenes: {path.name}")
    template_deck = Presentation(str(template))
    template_images = {
        hashlib.sha256(shape.image.blob).digest()
        for slide in template_deck.slides
        for shape in slide.shapes
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    }
    preserved_template_images = {
        hashlib.sha256(shape.image.blob).digest()
        for slide in list(template_deck.slides)[:2]
        for shape in slide.shapes
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    }
    for slide_number, slide in enumerate(deck.slides, 1):
        for shape in slide.shapes:
            if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                continue
            image_hash = hashlib.sha256(shape.image.blob).digest()
            if slide_number in (1, 2):
                if image_hash not in preserved_template_images:
                    raise AssertionError(f"La portada o agenda perdió la identidad visual: {path.name}")
            elif image_hash in template_images:
                raise AssertionError(f"Imagen sin adaptar de la plantilla en diapositiva {slide_number}: {path.name}")
            elif f"Grupo 09 · {profile['stack']} · {profile['protocol']}" not in (shape._element.nvPicPr.cNvPr.get("descr") or ""):
                raise AssertionError(f"Imagen no adaptada a la variante en diapositiva {slide_number}: {path.name}")
    with zipfile.ZipFile(path) as archive:
        bad = archive.testzip()
        if bad:
            raise AssertionError(f"ZIP inválido: {bad}")
        all_xml = "\n".join(archive.read(name).decode("utf-8", "ignore") for name in archive.namelist() if name.endswith(".xml"))
        forbidden = ("ANDRANGO CHRISTIAN", "CALAPAQUI GENESIS", "QUILUMBAQUIN JAIRO", "09/12/2024", "WSEurekaClient")
        stale = [marker for marker in forbidden if marker.casefold() in all_xml.casefold()]
        if stale:
            raise AssertionError(f"Texto heredado en {path.name}: {stale}")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", type=Path, required=True)
    parser.add_argument("--profile", choices=[item["key"] for item in PROFILES])
    args = parser.parse_args()
    root = args.root.resolve()
    template = root / TEMPLATE
    if not template.is_file():
        raise FileNotFoundError(f"Falta una de las dos referencias permitidas: {TEMPLATE.name}")
    profiles = [profile for profile in PROFILES if args.profile is None or profile["key"] == args.profile]
    for profile in profiles:
        output = update_deck(root, profile, template)
        print(f"{profile['key']}: {len(Presentation(str(output)).slides)} diapositivas adaptadas a la plantilla · {output.stat().st_size} bytes")


if __name__ == "__main__":
    main()

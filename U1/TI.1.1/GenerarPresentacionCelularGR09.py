"""
GenerarPresentacionCelularGR09.py
Genera PRESENTACION_CELULAR_RESTFUL_DOTNET_GR09.pptx para el proyecto
Comercializadora de Celulares, siguiendo el mismo patron que el generador
de EurekaBank (GenerarPresentacionesEurekaBankGR09.py).

Requiere: python-pptx, Pillow
  pip install python-pptx pillow
"""
from __future__ import annotations

import argparse
import os
import sys
import textwrap
from pathlib import Path

# ── dependencias opcionales ────────────────────────────────────────────────────
LIBS = Path(r"D:\tools\eureka-pptx-libs")
if LIBS.is_dir():
    sys.path.insert(0, str(LIBS))

try:
    from PIL import Image, ImageDraw, ImageFont, ImageOps
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE
    from pptx.util import Inches, Pt, Emu
    PIL_OK = True
except ImportError:
    PIL_OK = False

PROYECTO = "TI.1.1. PAO202651_CELULAR_BDDLOCAL_RESTFUL_NRC36884_GR09"

SLIDE_W  = 12192000   # EMU 16:9
SLIDE_H  =  6858000
EMU      = 914400

TEMA = {
    "fondo":   (15,  20,  35),   # azul muy oscuro
    "primario":(40, 100, 180),   # azul corporativo
    "acento":  (70, 200, 140),   # verde menta
    "texto":   (230, 235, 245),
    "subtexto":(160, 180, 210),
    "blanco":  (255, 255, 255),
}

DIAPOSITIVAS = [
    # (titulo, cuerpo, notas)
    ("COMERCIALIZADORA DE CELULARES",
     "GRUPO 09 · NRC 36884\nArquitectura de Software · .NET 10 · REST\n26 de octubre de 2026",
     "Portada"),

    ("CONTENIDO",
     "1. Objetivo\n2. Arquitectura REST\n3. Base de datos SQL Server\n"
     "4. API Comercializadora\n5. Cliente Consola\n6. Cliente WPF\n"
     "7. Cliente Web MVC\n8. Cliente MAUI Android\n9. Conclusiones",
     "Tabla de contenido"),

    ("1. OBJETIVO",
     "Implementar una comercializadora de celulares con arquitectura REST "
     "sobre .NET 10, que permita gestionar el catalogo, registrar ventas en "
     "efectivo y a credito directo mediante integracion con BanQuito, y "
     "consultar facturas y tablas de amortizacion desde cuatro clientes "
     "independientes.",
     "Objetivo"),

    ("2. ARQUITECTURA REST",
     "• API REST: ASP.NET Core · 127.0.0.1:19102\n"
     "• BanQuito (externo): 127.0.0.1:19101\n"
     "• 4 clientes: consola, WPF, MVC web, MAUI Android\n"
     "• Libreria compartida: Celulares.Cliente (DTOs + HTTP)\n"
     "• MVC aplicado en cada cliente independientemente",
     "Arquitectura"),

    ("DIAGRAMA DE ARQUITECTURA",
     "Los cuatro clientes consumen la API de la comercializadora por HTTP/JSON. "
     "La comercializadora consulta BanQuito solo para credito. "
     "Cada componente tiene su propia base o referencia remota.",
     "Arquitectura - diagrama"),

    ("3. BASE DE DATOS SQL SERVER",
     "Base: ComercializadoraCelulares en SQLEXPRESS\n\n"
     "CELULAR: ID_CELULAR, NOMBRE, PRECIO, ACTIVO\n"
     "FACTURA: ID_FACTURA, ID_OPERACION (UUID), ID_CELULAR FK,\n"
     "         FORMA_PAGO, TOTAL, ESTADO\n\n"
     "COD_CLIENTE_BANCO y COD_CREDITO son referencias remotas\n"
     "a BanQuitoCore (no claves foraneas locales).",
     "BDD"),

    ("MODELO ENTIDAD-RELACION FISICO",
     "CELULAR (1) ──< FACTURA (N)\n\n"
     "Scripts en 03. BDD:\n"
     "  01_CrearBases.sql · 02_DatosPrueba.sql\n\n"
     "Inicializar:\n"
     "  dotnet run --project 04.SERVIDOR\\Administracion -- --aplicar ...",
     "MER fisico"),

    ("4. API COMERCIALIZADORA",
     "GET  /api/celulares              → catalogo activo\n"
     "POST /api/celulares              → nuevo celular\n"
     "PUT  /api/celulares/{id}         → editar celular\n"
     "DELETE /api/celulares/{id}       → desactivar\n"
     "POST /api/facturas               → venta efectivo o credito\n"
     "GET  /api/facturas/{id}          → consultar factura\n"
     "GET  /api/facturas/{id}/amortiz  → tabla de cuotas",
     "API endpoints"),

    ("REGLAS DE NEGOCIO",
     "Efectivo: descuento 42 % sobre precio de lista\n\n"
     "Credito directo (via BanQuito):\n"
     "  · Cliente existente + deposito en ultimo mes\n"
     "  · Sin credito activo · Masculino >= 25 anos\n"
     "  · Tasa 16,5 % anual · Cuota fija · Plazo 3-18 meses\n\n"
     "Idempotencia: ID_OPERACION evita facturas duplicadas",
     "Reglas negocio"),

    ("5. CLIENTE CONSOLA",
     "Proyecto: Celulares.Consola (.NET 10)\n\n"
     "Separacion MVC:\n"
     "  · Vistas/ → presentacion en terminal\n"
     "  · Controladores/ → logica de interaccion\n"
     "  · Servicios (Celulares.Cliente) → HTTP JSON\n\n"
     "Ejecutar: dotnet run --project 04.CLICON\\Celulares.Consola",
     "Cliente consola"),

    ("CAPTURA CLIENTE CONSOLA",
     "Menu principal: catalogo, nueva venta, consultar factura, amortizacion.",
     "Captura consola"),

    ("6. CLIENTE ESCRITORIO WPF",
     "Proyecto: Celulares.Escritorio (WPF · .NET 10 · Windows)\n\n"
     "Separacion MVC:\n"
     "  · MainWindow.xaml → vista XAML\n"
     "  · Controladores/ControladorEscritorio.cs → logica\n"
     "  · Celulares.Cliente → transporte HTTP compartido\n\n"
     "Ejecutar: dotnet run --project 04.CLIESC\\Celulares.Escritorio",
     "Cliente WPF"),

    ("CAPTURA CLIENTE WPF",
     "Pantalla principal del escritorio: catalogo y formulario de venta.",
     "Captura WPF"),

    ("7. CLIENTE WEB MVC",
     "Proyecto: Celulares.Web (ASP.NET Core MVC · .NET 10)\n\n"
     "Separacion MVC nativa de ASP.NET:\n"
     "  · Views/ → Razor pages\n"
     "  · Controllers/TiendaController.cs → logica web\n"
     "  · Celulares.Cliente → servicios HTTP compartidos\n\n"
     "Ejecutar: dotnet run → http://127.0.0.1:19103",
     "Cliente Web"),

    ("CAPTURA CLIENTE WEB",
     "Catalogo de celulares en el navegador. Formulario de nueva venta.",
     "Captura Web"),

    ("8. CLIENTE MOVIL MAUI ANDROID",
     "Proyecto: Celulares.Movil (.NET MAUI · Android 12+)\n\n"
     "Separacion MVC:\n"
     "  · MainPage.cs → vista MAUI\n"
     "  · Controladores/ControladorMovil.cs → logica\n"
     "  · Celulares.Cliente → transporte HTTP compartido\n\n"
     "Emulador usa 10.0.2.2:19102 (no localhost)\n"
     "SDK: D:\\tools\\android-sdk (Android API 36)",
     "Cliente MAUI"),

    ("CAPTURA CLIENTE MOVIL",
     "Aplicacion MAUI Android en emulador GR10_Pixel_35: catalogo y venta.",
     "Captura MAUI"),

    ("DIAGRAMA DE CASOS DE USO",
     "Actor: Usuario (administrador / vendedor)\n\n"
     "· Gestionar catalogo (alta, edicion, baja logica)\n"
     "· Registrar venta en efectivo\n"
     "· Registrar venta a credito (incluye BanQuito)\n"
     "· Consultar factura\n"
     "· Consultar amortizacion",
     "Casos de uso"),

    ("DIAGRAMA DE CLASES",
     "Celular · Factura · ServicioCatalogo · ServicioFacturas\n"
     "ServicioBanQuito · ReglasFacturacion\n\n"
     "CelularesController · FacturasController (API REST)\n"
     "Celulares.Contratos: DTOs compartidos (solicitudes/respuestas)",
     "Clases"),

    ("DIAGRAMAS UML",
     "Power Designer (modelos .oom / .cdm / .ldm / .pdm):\n\n"
     "  01_CASOS_DE_USO/\n"
     "  02_ACTIVIDAD/\n"
     "  03_SECUENCIA/\n"
     "  04_CLASES/\n"
     "  05_ARQUITECTURA/\n\n"
     "Exportar PNG desde PowerDesigner y colocar en cada subcarpeta.",
     "UML"),

    ("9. CONCLUSIONES",
     "· La arquitectura REST con .NET 10 permite cuatro clientes con\n"
     "  codigo de transporte compartido y MVC independiente.\n\n"
     "· La idempotencia por ID_OPERACION garantiza consistencia ante\n"
     "  reintentos de red.\n\n"
     "· La integracion con BanQuito es desacoplada: solo HTTP; la\n"
     "  tienda no conoce el esquema bancario.\n\n"
     "· Pruebas automatizadas (9/9) verifican catalogo, facturas y credito.",
     "Conclusiones"),

    ("RECOMENDACIONES",
     "· Usar variables de entorno para endpoints y cadenas de conexion;\n"
     "  nunca en codigo ni capturas.\n\n"
     "· En produccion habilitar HTTPS y permisos minimos en SQLEXPRESS.\n\n"
     "· Para telefono fisico, iniciar la API con --urls http://0.0.0.0:19102\n"
     "  solo en redes confiables.\n\n"
     "· Mantener los modelos Power Designer sincronizados con el codigo.",
     "Recomendaciones"),

    ("BIBLIOGRAFÍA",
     "[1] Microsoft. ASP.NET Core documentation. https://docs.microsoft.com/aspnet/core\n"
     "[2] Microsoft. .NET MAUI documentation. https://docs.microsoft.com/dotnet/maui\n"
     "[3] Microsoft. Entity Framework Core. https://docs.microsoft.com/ef/core\n"
     "[4] SAP. PowerDesigner documentation. https://help.sap.com/powerdesigner",
     "Bibliografía"),
]


def make_slide_image(title: str, body: str, output: Path,
                     width: int = 1920, height: int = 1080,
                     asset: Path | None = None) -> Path:
    """Genera una imagen PNG para una diapositiva."""
    img  = Image.new("RGB", (width, height), TEMA["fondo"])
    draw = ImageDraw.Draw(img)

    # barra superior
    draw.rectangle((0, 0, width, 6), fill=TEMA["primario"])
    # barra inferior
    draw.rectangle((0, height - 6, width, height), fill=TEMA["acento"])
    # borde lateral izquierdo acento
    draw.rectangle((0, 0, 4, height), fill=TEMA["acento"])

    # fuentes
    font_t = _font("arialbd.ttf", 52)
    font_b = _font("arial.ttf",   30)
    font_f = _font("arial.ttf",   20)

    # titulo
    draw.text((60, 40), title, fill=TEMA["texto"], font=font_t)
    draw.rectangle((60, 100, 60 + min(len(title) * 28, width - 120), 104),
                   fill=TEMA["primario"])

    # imagen lateral si existe
    img_x = width
    if asset and asset.exists():
        try:
            with Image.open(asset) as src:
                src_rgb = src.convert("RGB")
                max_w, max_h = width // 2 - 40, height - 180
                contained = ImageOps.contain(src_rgb, (max_w, max_h),
                                             method=Image.Resampling.LANCZOS)
                img_x = width - contained.width - 30
                img_y = 130 + (height - 180 - contained.height) // 2
                img.paste(contained, (img_x, img_y))
        except Exception:
            pass

    # cuerpo de texto
    text_w = img_x - 80 if img_x < width else width - 80
    y = 130
    for line in body.split("\n"):
        wrapped = textwrap.wrap(line, width=max(30, text_w // 18)) or [""]
        for wline in wrapped:
            draw.text((60, y), wline, fill=TEMA["texto"], font=font_b)
            y += 40
        if not line.strip():
            y += 10

    # pie
    footer = "Comercializadora de Celulares  ·  Grupo 09  ·  NRC 36884"
    draw.text((60, height - 36), footer, fill=TEMA["subtexto"], font=font_f)

    img.save(output, format="PNG", optimize=True)
    return output


def _font(name: str, size: int):
    path = Path(r"C:\Windows\Fonts") / name
    try:
        return ImageFont.truetype(str(path), size)
    except Exception:
        return ImageFont.load_default()


def build_pptx(root: Path, output: Path) -> None:
    """Construye el PPTX directamente con python-pptx."""
    prs = Presentation()
    prs.slide_width  = Emu(SLIDE_W)
    prs.slide_height = Emu(SLIDE_H)

    blank_layout = prs.slide_layouts[6]  # completamente en blanco

    # Assets
    uml_dir  = root / "01. UML" / "PowerDesigner"
    mer_dir  = root / "02. MER" / "PowerDesigner"
    cap_dir  = root / "05. DOCUMENTACION" / "01. MANUAL TECNICO" / "Capturas"

    asset_map = {
        "DIAGRAMA DE ARQUITECTURA"    : uml_dir / "GR09_Celular_Arquitectura.png",
        "MODELO ENTIDAD-RELACION FISICO": mer_dir / "GR09_Celular_MER_Fisico.png",
        "CAPTURA CLIENTE CONSOLA"     : cap_dir / "Cliente_Consola.png",
        "CAPTURA CLIENTE WPF"         : cap_dir / "Cliente_WPF.png",
        "CAPTURA CLIENTE WEB"         : cap_dir / "Cliente_Web.png",
        "CAPTURA CLIENTE MOVIL"       : cap_dir / "Cliente_MAUI.png",
        "DIAGRAMA DE CASOS DE USO"    : uml_dir / "GR09_Celular_CasosUso.png",
        "DIAGRAMA DE CLASES"          : uml_dir / "GR09_Celular_Clases.png",
        "DIAGRAMAS UML"               : uml_dir / "GR09_Celular_Arquitectura.png",
    }

    import tempfile, os
    tmp = Path(tempfile.mkdtemp(prefix="celular-pptx-"))

    for idx, (title, body, _notes) in enumerate(DIAPOSITIVAS):
        slide = prs.slides.add_slide(blank_layout)

        # fondo negro/azul
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(*TEMA["fondo"])

        # generar imagen de la diapositiva
        asset = asset_map.get(title)
        slide_img = tmp / f"slide_{idx:02d}.png"
        make_slide_image(title, body, slide_img, asset=asset)

        # insertar imagen como fondo de slide
        slide.shapes.add_picture(
            str(slide_img),
            left=0, top=0,
            width=Emu(SLIDE_W), height=Emu(SLIDE_H)
        )

        # notas
        notes = slide.notes_slide
        notes.notes_text_frame.text = _notes

    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output))

    # limpiar tmp
    for f in tmp.iterdir():
        f.unlink()
    tmp.rmdir()

    print(f"Presentacion generada: {output}")
    print(f"Diapositivas: {len(DIAPOSITIVAS)}")


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Genera la presentacion PPTX del proyecto Celulares GR09")
    parser.add_argument("--root", type=Path, required=True,
                        help="Raiz del directorio TI.1.1")
    args = parser.parse_args()

    root_base = args.root.resolve()
    # buscar la carpeta del proyecto
    proyecto = root_base / PROYECTO
    if not proyecto.exists():
        # buscar con glob
        candidates = list(root_base.glob("*CELULAR*RESTFUL*"))
        if not candidates:
            raise FileNotFoundError(
                f"No se encontro el proyecto CELULAR en {root_base}")
        proyecto = candidates[0]

    if not PIL_OK:
        print("ERROR: Instale python-pptx y Pillow:")
        print("  pip install python-pptx pillow")
        sys.exit(1)

    output = (proyecto / "05. DOCUMENTACION" / "02. PRESENTACION"
              / "PRESENTACION_CELULAR_RESTFUL_DOTNET_GR09.pptx")
    build_pptx(proyecto, output)


if __name__ == "__main__":
    main()
